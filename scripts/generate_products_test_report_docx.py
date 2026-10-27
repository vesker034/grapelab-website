"""Генерация docs/GRL-15_отчет_тестирования_products.docx."""

from pathlib import Path

from docx import Document
from docx.shared import Pt

DOC_PATH = (
    Path(__file__).resolve().parent.parent
    / "docs"
    / "GRL-15_отчет_тестирования_products.docx"
)


def add_text(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    doc.add_paragraph()


def main() -> None:
    doc = Document()

    add_text(doc, "Отчёт по тестированию страницы /products", bold=True)
    add_text(doc, "Задача: GRL-15 — Проверка страницы продуктов /products")
    add_text(doc, "Проект: GrapeLab (grapelab-website)")
    add_text(doc, "Исполнитель тестирования: Андрей Толочек")
    add_text(doc, "Дата: 13.06.2026 (обновлённая версия отчёта)")

    add_text(doc, "Важно: соответствие чеклиста и кода", bold=True)
    add_text(
        doc,
        "В чеклисте GRL-15 сущности названы Product / ProductStatus / ProductCategory. "
        "В репозитории админка реализована под именами Platform / PlatformStatus / PlatformCategory "
        "на ветке origin/koptev/ecosystem-page (автор — Коптев). "
        "На origin/main есть только публичная /products с mock-данными; админка продуктов в main не смержена.",
    )

    add_text(doc, "Проверенные ветки", bold=True)
    add_table(
        doc,
        ["Ветка", "Коммит", "Что есть"],
        [
            [
                "origin/main",
                "d8a56c5",
                "Публичная /products (mockProducts). Админки Product/Platform нет.",
            ],
            [
                "origin/koptev/ecosystem-page",
                "18753b4",
                "/products из БД (Platform), админка Platform*, страница /platforms.",
            ],
        ],
    )

    add_text(doc, "Условия тестирования", bold=True)
    for line in [
        "• Репозиторий: grapelab-website (code.grapelab.ru/shashkov/grapelab-website).",
        "• git fetch --all по всем remote-веткам; поиск Product и Platform в schema/admin.",
        "• Часть 1 (main): npm run build, npm run dev, HTTP /ru/products, /ar/products.",
        "• Часть 1–2 (ecosystem-page): анализ кода (runtime-прогон админки без staff-аккаунта не выполнялся).",
        "• Прод grapelab.ru не тестировался.",
    ]:
        add_text(doc, line)

    add_text(doc, "Сводный результат", bold=True)
    add_table(
        doc,
        ["Часть", "origin/main", "origin/koptev/ecosystem-page"],
        [
            ["1 — публичная /products", "~85%, с замечаниями", "~90%, с замечаниями (данные из БД)"],
            ["2 — админка", "❌ отсутствует", "✅ реализована (Platform*), по коду"],
        ],
    )

    # --- PART 1 MAIN ---
    add_text(doc, "Часть 1. Публичная /products — origin/main", bold=True)

    add_text(doc, "1.1. Hero-секция", bold=True)
    add_table(
        doc,
        ["№", "Проверка", "Результат", "Комментарий"],
        [
            ["1", "Заголовок, подзаголовок, бейдж, CTA", "✅ OK", "ProductsHero"],
            ["2", "Метрики с иконками", "✅ OK", "productHeroMetrics"],
            ["3", "Анимации, пульсация точки", "✅ OK", "framer-motion, animate-pulse"],
            ["4", "Вращение орбиты", "✅ OK", "orbitSpin в globals.css"],
            ["5", "Орбита desktop / скрыта mobile", "✅ OK", "hidden lg:flex"],
        ],
    )

    add_text(doc, "1.2. Каталог продуктов", bold=True)
    add_table(
        doc,
        ["№", "Проверка", "Результат", "Комментарий"],
        [
            ["1", "Поиск по названию и описанию", "✅ OK", ""],
            ["2", "Фильтры: категории, статусы, платформы", "✅ OK", ""],
            ["3", "Кнопка сброса при активных фильтрах", "✅ OK", ""],
            ["4", "Сортировка", "⚠️ Частично", "«По популярности» не сортирует"],
            ["5", "Пагинация 6 + скролл к сетке", "✅ OK", ""],
            ["6", "Mobile panel + затемнение", "✅ OK", ""],
            ["7", "Пустой результат + сброс", "✅ OK", ""],
        ],
    )

    add_text(doc, "1.3–1.5. Карточка, локализация, адаптив (main)", bold=True)
    add_table(
        doc,
        ["Блок", "Результат", "Комментарий"],
        [
            ["1.3 Карточка", "⚠️", "Баг метрик: сырые ключи activeUsers/completedCourses; hover OK"],
            ["1.4 Локализация", "⚠️", "UI ru/en/ar OK; названия из mock (англ.)"],
            ["1.5 Адаптив", "✅ OK", "Mobile panel, desktop sidebar; одна колонка карточек"],
        ],
    )

    # --- PART 1 ECOSYSTEM ---
    add_text(doc, "Часть 1. Публичная /products — origin/koptev/ecosystem-page", bold=True)
    add_text(
        doc,
        "На этой ветке /products загружает данные из БД: getPlatformCatalogData() → mapPlatformToProduct(). "
        "Названия и описания берутся из локализаций Platform (platforms_name, platforms_description).",
    )
    add_table(
        doc,
        ["№", "Проверка", "Результат", "Комментарий"],
        [
            ["1.1 Hero", "✅ OK", "Тот же ProductsHero"],
            ["1.2 Каталог", "⚠️", "Те же замечания: sort popular, одна колонка"],
            ["1.3 Карточка", "⚠️", "t(metric.label) — баг сохраняется; imagePath из БД не используется в ProductCard (иконка по category)"],
            ["1.4 Локализация", "✅ OK", "Имена из БД/локализаций; RTL ar поддержан"],
            ["1.5 Адаптив", "✅ OK", "Без изменений относительно main"],
        ],
    )
    add_text(
        doc,
        "Цвет статуса: на странице /platforms берётся из PlatformStatus.color (#RRGGBB), "
        "но на /products в ProductCard — захардкоженный statusColorMap (цвет из админки не применяется).",
    )

    # --- PART 2 ---
    add_text(doc, "Часть 2. Админка — origin/main", bold=True)
    add_text(doc, "CRUD Product / ProductStatus / ProductCategory на main отсутствует. Раздел не тестировался.")

    add_text(doc, "Часть 2. Админка — origin/koptev/ecosystem-page", bold=True)
    add_text(
        doc,
        "Соответствие чеклиста → реализация в коде:",
    )
    add_table(
        doc,
        ["Пункт чеклиста", "В коде", "URL админки"],
        [
            ["Product CRUD", "Platform", "/admin/Platform"],
            ["ProductStatus", "PlatformStatus (+ color HEX)", "/admin/PlatformStatus"],
            ["ProductCategory", "PlatformCategory (+ icon upload)", "/admin/PlatformCategory"],
        ],
    )

    add_text(doc, "2.1. CRUD продуктов (Platform)", bold=True)
    add_table(
        doc,
        ["№", "Проверка", "Результат", "Комментарий"],
        [
            ["1", "Создание со всеми полями", "✅ по коду", "Prisma model Platform, NextAdmin options.ts"],
            ["2", "Редактирование / сохранение", "✅ по коду", "LocalizedFieldsInput для rawName + поля формы"],
            ["3", "Удаление", "✅ по коду", "Стандартный NextAdmin CRUD"],
            ["4", "Список: колонки, поиск, связи", "✅ по коду", "search rawName; formatters category/status"],
        ],
    )

    add_text(doc, "2.2. Кастомные поля", bold=True)
    add_table(
        doc,
        ["№", "Проверка", "Результат", "Комментарий"],
        [
            ["1", "Платформы (web, ios, …)", "✅ по коду", "platforms String[]; валидация allowed list"],
            ["2", "Метрики: add/remove пар", "✅ по коду", "AspectsListEditor + JSON schema metrics"],
            ["3", "Загрузка иконки", "✅ по коду", "imagePath → ImageUploadInput; scope platforms/platformcategories"],
        ],
    )

    add_text(doc, "2.3. Валидация", bold=True)
    add_table(
        doc,
        ["Поле", "Чеклист", "Результат", "Комментарий"],
        [
            ["rawName", "латиница, цифры, _", "✅", "slugFormat + maxLength 100"],
            ["name", "2–200, обязательно", "⚠️", "Через локализацию platforms_name (LocalizedFieldsInput), не отдельное поле name"],
            ["status", "обязателен", "✅", "required('Статус')"],
            ["category", "обязателен", "⚠️", "categoryId nullable в schema; явного required нет"],
            ["url", "валидный URL", "✅", "urlFormat (опционально)"],
            ["description", "≤ 2000", "⚠️", "В локализации platforms_description; отдельного max 2000 в options не найдено"],
            ["sortOrder", "≥ 0", "✅", "nonNegativeInt"],
            ["platforms, metrics", "корректный JSON/массив", "✅", "validate + validateJsonBySchema"],
        ],
    )

    add_text(doc, "2.4. Локализационные хуки", bold=True)
    add_table(
        doc,
        ["№", "Проверка", "Результат", "Комментарий"],
        [
            ["1", "Запись в Localization после save", "⚠️", "localizationConfig: Platform, PlatformStatus, PlatformCategory; клиент LocalizedFieldsInput → /api/admin/localizations"],
            ["2", "afterDb hook для Platform", "⚠️", "saveLocalizations — заглушка (console.log); как у Service на main"],
            ["3", "Обновление фронта после save", "✅ по коду", "/products читает локализации через getPlatformCatalogData"],
        ],
    )

    add_text(doc, "2.5. Связанные модели", bold=True)
    add_table(
        doc,
        ["№", "Проверка", "Результат", "Комментарий"],
        [
            ["1", "PlatformStatus CRUD + color #RRGGBB", "✅ по коду", "HEX-валидация в options.ts"],
            ["2", "PlatformCategory CRUD", "✅ по коду", "icon, sortOrder, isVisible"],
            ["3", "Цвет статуса на фронте /products", "❌", "ProductCard — hardcoded map; цвет из админки только на /platforms"],
        ],
    )
    add_text(
        doc,
        "Замечание: Platform / PlatformStatus / PlatformCategory не добавлены в быстрые ссылки AdminDashboard по умолчанию, "
        "но доступны напрямую по URL (/admin/Platform и т.д.).",
    )

    add_text(doc, "Выявленные дефекты и расхождения", bold=True)
    add_table(
        doc,
        ["№", "Приоритет", "Описание", "Где"],
        [
            [
                "1",
                "Средний",
                "Метрики карточки: сырые ключи вместо перевода (t(metric.label))",
                "ProductCard.tsx — обе ветки",
            ],
            [
                "2",
                "Низкий",
                "Сортировка «По популярности» не реализована",
                "ProductsCatalog.tsx",
            ],
            [
                "3",
                "Низкий",
                "Цвет PlatformStatus из админки не используется на /products",
                "ecosystem-page: ProductCard vs PlatformCard",
            ],
            [
                "4",
                "Низкий",
                "/products нет в sitemap.ts",
                "main",
            ],
            [
                "5",
                "Процесс",
                "Админка Platform* не смержена в main/dev",
                "koptev/ecosystem-page",
            ],
        ],
    )

    add_text(doc, "Рекомендации", bold=True)
    for line in [
        "• Смержить origin/koptev/ecosystem-page в main/dev — иначе на prod админка продуктов недоступна.",
        "• Исправить ProductCard: метрики и цвет статуса из БД (как на /platforms).",
        "• Добавить case «popular» в сортировку или убрать пункт из UI.",
        "• Добавить Platform* в AdminDashboard и /products в sitemap.",
        "• После merge — повторить smoke-тест админки с staff-аккаунтом (create/edit/delete).",
    ]:
        add_text(doc, line)

    add_text(doc, "Заключение", bold=True)
    add_text(
        doc,
        "Публичная страница /products реализована (main — mock, ecosystem-page — БД). "
        "Админка из части 2 чеклиста существует на ветке koptev/ecosystem-page под именем Platform*, "
        "но не входит в main. Первоначальный отчёт «админка 0%» относился только к main; "
        "после поиска по всем веткам админка найдена и проанализирована по коду. "
        "Рекомендуется merge ветки и повторный runtime-тест админки перед закрытием GRL-15.",
    )

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")


if __name__ == "__main__":
    main()
