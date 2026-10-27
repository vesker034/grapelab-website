"""Обновление DOCX: контекст edu, внутренние ЛК, экосистема GrapeLab."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

DOC_PATH = (
    Path(__file__).resolve().parent.parent
    / "docs"
    / "GrapeLab-OAuth-SSO-issledovanie.docx"
)


def set_text(para: Paragraph, text: str) -> None:
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def insert_after(para: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = deepcopy(para._p)
    para._p.addnext(new_p)
    new_para = Paragraph(new_p, para._parent)
    if style:
        new_para.style = style
    set_text(new_para, text)
    return new_para


def find_para(doc: Document, prefix: str) -> Paragraph | None:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    return None


def replace_cell_text(doc: Document, old: str, new: str) -> None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if old in para.text:
                        set_text(para, para.text.replace(old, new))


def para_index(doc: Document, para: Paragraph) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p._p is para._p:
            return i
    raise ValueError("paragraph not found")


def main() -> None:
    doc = Document(DOC_PATH)

    note = find_para(doc, "Примечание.")
    if note:
        set_text(
            note,
            "Примечание. Документ описывает исследование единой авторизации для экосистемы "
            "GrapeLab (вуз + компания + внутренние сервисы). Краткие выводы — в разделе 12.",
        )

    terms = find_para(doc, "1. Термины")
    if terms:
        ti = para_index(doc, terms)
        prev = doc.paragraphs[ti - 1] if ti > 0 else terms
        if not prev.text.startswith("0. Контекст"):
            ctx = insert_after(prev, "0. Контекст организации GrapeLab")
            p = insert_after(
                ctx,
                "GrapeLab — студенческая IT-компания при вузе. В команде в основном студенты. "
                "Экосистема включает несколько типов сервисов с разными личными кабинетами:",
            )
            p = insert_after(
                p,
                "edu — официальный сайт вуза: курсы по предметам, учебные материалы, ЛК студентов и преподавателей.",
            )
            p = insert_after(
                p,
                "Внутренние сервисы GrapeLab — сайты и инструменты для работников компании (в т.ч. доступ к git, задачи, внутренние порталы); у каждого есть ЛК работника.",
            )
            p = insert_after(
                p,
                "grapelab.ru — публичный сайт компании (услуги, кейсы, блог). В новой версии (v2) публичного ЛК нет; есть админка для сотрудников (staff), редактирующих контент.",
            )
            p = insert_after(
                p,
                "Проблема: один человек (студент и работник GrapeLab) может иметь отдельные учётные записи "
                "на edu, на внутренних сервисах и в админке — разные пароли и базы пользователей.",
            )
            insert_after(
                p,
                "Цель SSO: один аккаунт на auth.grapelab.ru, разные роли (student, teacher, worker, staff) "
                "определяют доступ на edu, внутренние сервисы и grapelab.ru.",
            )

    apply_para = find_para(doc, "Применимость для GrapeLab:")
    if apply_para:
        set_text(
            apply_para,
            "Применимость для GrapeLab: основной поток для edu, внутренних сервисов, grapelab.ru и админки.",
        )

    # Section 5 rewrite
    s51 = find_para(doc, "5.1.")
    if s51:
        set_text(s51, "5.1. Публичный сайт grapelab.ru (v2)")
        idx = para_index(doc, s51)
        if idx + 1 < len(doc.paragraphs):
            set_text(
                doc.paragraphs[idx + 1],
                "Маркетинговый сайт компании. Публичного ЛК для посетителей пока нет.",
            )
        if idx + 2 < len(doc.paragraphs):
            set_text(
                doc.paragraphs[idx + 2],
                "Админка контента: NextAuth + Credentials (email/пароль), только isStaff. JWT-сессия 8 часов.",
            )
        if idx + 3 < len(doc.paragraphs) and "JWT" in doc.paragraphs[idx + 3].text:
            set_text(doc.paragraphs[idx + 3], "")

    s52 = find_para(doc, "5.2.")
    if s52:
        set_text(s52, "5.2. Сайт вуза (edu)")
        idx = para_index(doc, s52)
        if idx + 1 < len(doc.paragraphs):
            set_text(
                doc.paragraphs[idx + 1],
                "Платформа вуза с курсами по предметам. Есть ЛК студентов и преподавателей, своя база пользователей.",
            )
        if idx + 2 < len(doc.paragraphs):
            p = doc.paragraphs[idx + 2]
            set_text(
                p,
                "Основной кандидат на первое подключение к центральному IdP (наибольшее число пользователей).",
            )
            if not find_para(doc, "5.3. Внутренние"):
                h53 = insert_after(p, "5.3. Внутренние сервисы GrapeLab")
                p = insert_after(
                    h53,
                    "Сайты и инструменты для работников компании: личные кабинеты, доступ к git, внутренние задачи. "
                    "Учётные записи работников уже существуют — без них невозможна повседневная работа в GrapeLab.",
                )
                insert_after(
                    p,
                    "При внедрении SSO необходимо аудировать, сколько отдельных систем входа используется сейчас "
                    "и где хранятся учётки работников (одна БД или несколько).",
                )

    arch = find_para(doc, "Центральный IdP на auth.grapelab.ru")
    if arch:
        set_text(
            arch,
            "Центральный IdP на auth.grapelab.ru (собственные серверы GrapeLab). Все сервисы экосистемы — "
            "OIDC clients: edu, внутренние порталы, grapelab.ru. Единая регистрация и вход на IdP; "
            "каждый сервис хранит только привязку sub → локальный профиль и проверяет роли.",
        )

    scheme = find_para(doc, "auth.grapelab.ru (IdP: Authentik)")
    if scheme:
        set_text(
            scheme,
            "auth.grapelab.ru (IdP: Authentik)\n"
            "       │\n"
            "       ├── edu.grapelab.ru / edu.*     — client_id: grapelab-edu (роли: student, teacher)\n"
            "       ├── внутренние сервисы          — client_id: grapelab-internal (роль: worker)\n"
            "       ├── grapelab.ru                 — client_id: grapelab-web\n"
            "       └── админка grapelab.ru         — client_id: grapelab-admin (роль: staff)",
        )

    roles = find_para(doc, "Роли: groups в Authentik")
    if roles:
        set_text(
            roles,
            "Роли в Authentik: student, teacher, worker, staff — один пользователь может иметь несколько ролей; "
            "синхронизация при входе.",
        )

    clients = find_para(doc, "Отдельные client_id: web, edu, admin")
    if clients:
        set_text(
            clients,
            "Отдельные client_id: edu, internal, web, admin — изоляция redirect URI и политик доступа.",
        )

    mig_title = find_para(doc, "10. Миграция со старого prod")
    if mig_title:
        set_text(mig_title, "10. Миграция существующих учётных записей")
        idx = para_index(doc, mig_title)
        replacements = [
            (
                idx + 1,
                "Инвентаризация: edu (студенты/преподаватели), внутренние сервисы (работники), админка grapelab.ru (staff).",
            ),
            (
                idx + 2,
                "Сопоставление пользователей по email: один человек — одна запись в IdP (студент GrapeLab часто = студент вуза).",
            ),
            (
                idx + 3,
                "Импорт в Authentik (CSV/API); перенос хешей паролей или принудительный сброс — по результатам аудита.",
            ),
            (
                idx + 4,
                "Параллельный период: старый и новый вход на каждом сервисе (флаг переключения).",
            ),
            (
                idx + 5,
                "Поэтапное отключение legacy-auth: сначала PoC на staging, затем edu или внутренний сервис, затем остальные.",
            ),
        ]
        for i, text in replacements:
            if i < len(doc.paragraphs):
                set_text(doc.paragraphs[i], text)

    concl = find_para(doc, "Следующий шаг:")
    if concl:
        set_text(
            concl,
            "Следующий шаг: staging IdP + PoC; аудит учётных записей edu и внутренних сервисов; "
            "согласование с командой количества отдельных систем входа.",
        )

    doc.save(DOC_PATH)
    print(f"Updated: {DOC_PATH}")


if __name__ == "__main__":
    main()
